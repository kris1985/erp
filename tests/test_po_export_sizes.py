"""采购单导出：有码数合并为一行，数量含损耗。"""

from io import BytesIO

from openpyxl import load_workbook

from app.services.purchase_order_export import build_purchase_order_workbook


def _detail(*, status: str) -> dict:
    lines = [
        {
            "supplier_product_id": 1,
            "supplier_product_code": "ML-1",
            "supplier_product_name": "面料",
            "pricing_unit_name": "米",
            "size_value": "40",
            "size_sort_order": 40,
            "qty": 6,
            "loss_qty": 1,
            "unit_price": 2,
            "order_no": "SO-1",
        },
        {
            "supplier_product_id": 1,
            "supplier_product_code": "ML-1",
            "supplier_product_name": "面料",
            "pricing_unit_name": "米",
            "size_value": "42",
            "size_sort_order": 42,
            "qty": 4,
            "loss_qty": 0,
            "unit_price": 2,
            "order_no": "SO-1",
        },
        {
            "supplier_product_id": 2,
            "supplier_product_code": "ZX-1",
            "supplier_product_name": "纸箱",
            "pricing_unit_name": "个",
            "size_value": None,
            "qty": 3,
            "loss_qty": 1,
            "unit_price": 5,
            "order_no": "SO-1",
        },
    ]
    return {
        "status": status,
        "po_no": "PO-1",
        "buyer_name": "厂",
        "lines": lines,
        "summary_lines": [],
    }


def _sheet_values(ws):
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def test_draft_export_merges_sizes_and_adds_loss():
    raw = build_purchase_order_workbook(_detail(status="draft"), include_internal=True)
    wb = load_workbook(BytesIO(raw))
    supplier = _sheet_values(wb["供应商联"])
    flat = [cell for row in supplier for cell in row]
    assert "码数" in flat
    assert "40" in flat and "42" in flat
    # 40 码 6+1，42 码 4，纸箱 3+1；面料金额 (7+4)*2=22，纸箱 4*5=20
    data = [row for row in supplier if row and row[0] in ("ML-1", "ZX-1")]
    assert len(data) == 2
    fabric = next(row for row in data if row[0] == "ML-1")
    box = next(row for row in data if row[0] == "ZX-1")
    assert 7 in fabric
    assert 4 in fabric
    assert 11 in fabric
    assert 22 in fabric or 22.0 in fabric
    assert box[3] == "—"
    assert 4 in box

    internal = _sheet_values(wb["内部联"])
    internal_rows = [row for row in internal if row and row[0] == "ML-1"]
    assert len(internal_rows) == 1
    assert 7 in internal_rows[0] and 4 in internal_rows[0]


def test_ordered_export_does_not_add_loss_again():
    raw = build_purchase_order_workbook(_detail(status="ordered"))
    wb = load_workbook(BytesIO(raw))
    supplier = _sheet_values(wb["供应商联"])
    fabric = next(row for row in supplier if row and row[0] == "ML-1")
    # 已下单数量已含损耗，不再加：40 仍是 6，合计 10
    assert 7 not in fabric
    assert 6 in fabric
    assert 10 in fabric
