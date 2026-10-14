"""物料色卡 Excel 导入。"""

from __future__ import annotations

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Color, MaterialCategory, Partner, PricingUnit, SupplierProduct

MAX_IMPORT_ROWS = 2000
MAX_ERRORS = 50

HEADERS = ("物料编号", "名称", "分类", "颜色", "单价", "计价单位", "供应商")

_HEADER_FIELDS = {
    "物料编号": "code",
    "编号": "code",
    "商品编号": "code",
    "名称": "name",
    "物料名称": "name",
    "分类": "category",
    "物料分类": "category",
    "颜色": "color",
    "物料颜色": "color",
    "单价": "price",
    "计价单位": "unit",
    "单位": "unit",
    "供应商": "supplier",
}


class SupplierProductImportError(Exception):
    def __init__(self, message: str):
        self.message = message
        super().__init__(message)


def build_supplier_product_import_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "物料色卡"
    ws.append(list(HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    widths = (18, 24, 14, 12, 12, 14, 20)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.freeze_panes = "A2"
    note = wb.create_sheet("说明")
    note["A1"] = "填写说明"
    note["A1"].font = Font(bold=True)
    note["A2"] = "在「物料色卡」表填写，不要改表头。"
    note["A3"] = "物料编号、供应商必填。"
    note["A4"] = "分类、颜色、计价单位须与系统里已有的名称完全一致，没有的请先在基础资料中新增。"
    note["A5"] = "物料编号已存在时按表格更新；留空的单元格不会覆盖原来的值。"
    note["A6"] = "图片不能通过表格导入。"
    note.column_dimensions["A"].width = 88
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_supplier_products_xlsx(db: Session, tenant_id: int, content: bytes) -> dict:
    rows = _read_rows(content)
    header_at, columns = _locate_header(rows)
    suppliers_on, suppliers_off = _named_index(
        db.scalars(
            select(Partner).where(Partner.tenant_id == tenant_id, Partner.is_supplier.is_(True))
        ).all()
    )
    categories_on, categories_off = _named_index(
        db.scalars(select(MaterialCategory).where(MaterialCategory.tenant_id == tenant_id)).all()
    )
    units_on, units_off = _named_index(
        db.scalars(select(PricingUnit).where(PricingUnit.tenant_id == tenant_id)).all()
    )
    colors = {
        (c.name or "").strip(): c
        for c in db.scalars(select(Color).where(Color.tenant_id == tenant_id)).all()
        if (c.name or "").strip()
    }
    existing = {
        p.product_code: p
        for p in db.scalars(
            select(SupplierProduct).where(SupplierProduct.tenant_id == tenant_id)
        ).all()
    }

    created = 0
    updated = 0
    errors: list[dict] = []
    seen: set[str] = set()
    data_rows = 0
    for offset, raw in enumerate(rows[header_at + 1 :], start=header_at + 2):
        parsed = _parse_row(raw, columns)
        if parsed is None:
            continue
        if isinstance(parsed, str):
            data_rows += 1
            errors.append({"row": offset, "message": parsed})
            continue
        data_rows += 1
        if data_rows > MAX_IMPORT_ROWS:
            raise SupplierProductImportError(f"一次最多导入 {MAX_IMPORT_ROWS} 行")
        code = parsed["code"]
        if code in seen:
            errors.append({"row": offset, "message": f"物料编号「{code}」在表格中重复"})
            continue
        seen.add(code)
        message = _apply_row(
            db,
            tenant_id,
            parsed,
            existing,
            suppliers_on,
            suppliers_off,
            categories_on,
            categories_off,
            colors,
            units_on,
            units_off,
        )
        if message:
            errors.append({"row": offset, "message": message})
            continue
        if parsed["existed"]:
            updated += 1
        else:
            created += 1

    if data_rows == 0:
        raise SupplierProductImportError("没有可导入的数据")
    if created or updated:
        db.commit()
    return {
        "created": created,
        "updated": updated,
        "failed": len(errors),
        "errors": errors[:MAX_ERRORS],
    }


def _read_rows(content: bytes) -> list[tuple]:
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except (InvalidFileException, BadZipFile, ValueError, OSError) as exc:
        raise SupplierProductImportError("无法读取 Excel，请使用 .xlsx 模板") from exc
    try:
        ws = wb["物料色卡"] if "物料色卡" in wb.sheetnames else wb.active
        if ws is None:
            raise SupplierProductImportError("Excel 里没有工作表")
        return [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _locate_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for index, raw in enumerate(rows[:8]):
        columns: dict[str, int] = {}
        for col, value in enumerate(raw):
            field = _HEADER_FIELDS.get(_norm_header(value))
            if field and field not in columns:
                columns[field] = col
        if "code" in columns and "supplier" in columns:
            return index, columns
    raise SupplierProductImportError("请使用物料色卡导入模板，表头需包含物料编号和供应商")


def _parse_row(raw: tuple, columns: dict[str, int]) -> dict | str | None:
    def cell(field: str) -> str:
        idx = columns.get(field)
        if idx is None or idx >= len(raw):
            return ""
        return _cell_text(raw[idx])

    code = cell("code")
    name = cell("name")
    category = cell("category")
    color = cell("color")
    unit = cell("unit")
    supplier = cell("supplier")
    price_text = cell("price")
    price_raw = raw[columns["price"]] if "price" in columns and columns["price"] < len(raw) else None
    if not any((code, name, category, color, price_text, unit, supplier)):
        return None
    if not code:
        return "请填写物料编号"
    if len(code) > 50:
        return "物料编号不能超过 50 字"
    if len(name) > 100:
        return "名称不能超过 100 字"
    if not supplier:
        return "请填写供应商"
    price_set = price_text != "" or (price_raw is not None and price_raw != "")
    price = None
    if price_set:
        try:
            price = _parse_price(price_raw if price_text == "" else price_text)
        except SupplierProductImportError as exc:
            return exc.message
        if price is None:
            price_set = False
    return {
        "code": code,
        "name": name or None,
        "category": category or None,
        "color": color or None,
        "unit": unit or None,
        "supplier": supplier,
        "price": price,
        "price_set": price_set,
        "existed": False,
    }


def _apply_row(
    db: Session,
    tenant_id: int,
    parsed: dict,
    existing: dict[str, SupplierProduct],
    suppliers_on: dict[str, Partner],
    suppliers_off: dict[str, Partner],
    categories_on: dict[str, MaterialCategory],
    categories_off: dict[str, MaterialCategory],
    colors: dict[str, Color],
    units_on: dict[str, PricingUnit],
    units_off: dict[str, PricingUnit],
) -> str | None:
    supplier, message = _resolve_named(parsed["supplier"], suppliers_on, suppliers_off, "供应商")
    if message or supplier is None:
        return message or "供应商不存在"
    category = None
    if parsed["category"]:
        category, message = _resolve_named(parsed["category"], categories_on, categories_off, "分类")
        if message:
            return message
    color = None
    if parsed["color"]:
        color = colors.get(parsed["color"])
        if color is None:
            return "颜色不存在"
    unit = None
    if parsed["unit"]:
        unit, message = _resolve_named(parsed["unit"], units_on, units_off, "计价单位")
        if message:
            return message

    product = existing.get(parsed["code"])
    if product is None:
        product = SupplierProduct(
            tenant_id=tenant_id,
            product_code=parsed["code"],
            name=parsed["name"],
            category_id=category.id if category else None,
            pricing_unit_id=unit.id if unit else None,
            unit_price=parsed["price"] if parsed["price_set"] else None,
            color_id=color.id if color else None,
            partner_id=supplier.id,
            is_active=True,
        )
        db.add(product)
        existing[parsed["code"]] = product
        return None

    parsed["existed"] = True
    if parsed["name"]:
        product.name = parsed["name"]
    if category is not None:
        product.category_id = category.id
    if color is not None:
        product.color_id = color.id
    if unit is not None:
        product.pricing_unit_id = unit.id
    if parsed["price_set"]:
        product.unit_price = parsed["price"]
    product.partner_id = supplier.id
    return None


def _named_index(rows) -> tuple[dict, dict]:
    active: dict = {}
    inactive: dict = {}
    for row in rows:
        key = (row.name or "").strip()
        if not key:
            continue
        if getattr(row, "is_active", True):
            active[key] = row
        else:
            inactive.setdefault(key, row)
    return active, inactive


def _resolve_named(name: str, active: dict, inactive: dict, label: str):
    found = active.get(name)
    if found is not None:
        return found, None
    if name in inactive:
        return None, f"{label}已停用"
    return None, f"{label}不存在"


def _norm_header(value) -> str:
    text = _cell_text(value).replace(" ", "").replace("\u3000", "")
    text = text.rstrip("*＊")
    for suffix in ("（必填）", "(必填)", "（选填）", "(选填)"):
        if text.endswith(suffix):
            text = text[: -len(suffix)]
    return text


def _cell_text(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def _parse_price(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise SupplierProductImportError("单价不是数字")
    text = _cell_text(value).replace(",", "").replace("，", "")
    if not text:
        return None
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise SupplierProductImportError("单价不是数字") from exc
    if amount < 0:
        raise SupplierProductImportError("单价不能为负")
    return amount.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
