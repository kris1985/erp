"""产品开发物料明细 Excel 解析。只解析，不写入产品。"""

from __future__ import annotations

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Color, Department, MaterialCategory, Partner, PricingUnit, ProcessSegment, SupplierProduct

MAX_IMPORT_ROWS = 2000
MAX_ERRORS = 50

HEADERS = ("物料编号", "用量", "消耗部门")

_HEADER_FIELDS = {
    "物料编号": "code",
    "编号": "code",
    "用量": "qty",
    "数量": "qty",
    "消耗部门": "department",
    "部门": "department",
}


class OwnProductMaterialImportError(Exception):
    def __init__(self, message: str):
        self.message = message
        super().__init__(message)


def build_own_product_material_import_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "产品物料"
    ws.append(list(HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for idx, width in enumerate((18, 12, 16), start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.freeze_panes = "A2"
    note = wb.create_sheet("说明")
    note["A1"] = "填写说明"
    note["A1"].font = Font(bold=True)
    note["A2"] = "在「产品物料」表填写，不要改表头。"
    note["A3"] = "导入后记入当前正在编辑的成品颜色。"
    note["A4"] = "物料编号须是物料色卡里的编号。"
    note["A5"] = "用量留空时，新物料按 1；已经在表里的物料不改用量。"
    note["A6"] = "消耗部门留空时，新物料用分类的默认消耗部门；已经在表里的物料不改部门。"
    note["A7"] = "物料编号不能重复。"
    note.column_dimensions["A"].width = 88
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_own_product_materials_xlsx(db: Session, tenant_id: int, content: bytes) -> dict:
    rows = _read_rows(content)
    header_at, columns = _locate_header(rows)
    products = {
        p.product_code: p
        for p in db.scalars(
            select(SupplierProduct).where(SupplierProduct.tenant_id == tenant_id)
        ).all()
    }
    categories = {
        c.id: c
        for c in db.scalars(select(MaterialCategory).where(MaterialCategory.tenant_id == tenant_id)).all()
    }
    segments = list(
        db.scalars(
            select(ProcessSegment)
            .where(ProcessSegment.tenant_id == tenant_id)
            .order_by(ProcessSegment.sort_order, ProcessSegment.id)
        ).all()
    )
    departments = list(
        db.scalars(select(Department).where(Department.tenant_id == tenant_id)).all()
    )
    partner_ids = {p.partner_id for p in products.values()}
    color_ids = {p.color_id for p in products.values() if p.color_id}
    unit_ids = {p.pricing_unit_id for p in products.values() if p.pricing_unit_id}
    partners = {
        p.id: p
        for p in db.scalars(select(Partner).where(Partner.id.in_(partner_ids))).all()
    } if partner_ids else {}
    colors = {
        c.id: c
        for c in db.scalars(select(Color).where(Color.id.in_(color_ids))).all()
    } if color_ids else {}
    units = {
        u.id: u
        for u in db.scalars(select(PricingUnit).where(PricingUnit.id.in_(unit_ids))).all()
    } if unit_ids else {}

    items: list[dict] = []
    errors: list[dict] = []
    seen: set[str] = set()
    data_rows = 0
    for offset, raw in enumerate(rows[header_at + 1 :], start=header_at + 2):
        parsed = _parse_row(raw, columns)
        if parsed is None:
            continue
        data_rows += 1
        if data_rows > MAX_IMPORT_ROWS:
            raise OwnProductMaterialImportError(f"一次最多导入 {MAX_IMPORT_ROWS} 行")
        if isinstance(parsed, str):
            errors.append({"row": offset, "message": parsed})
            continue
        product = products.get(parsed["code"])
        if product is None:
            errors.append({"row": offset, "message": "物料不存在"})
            continue
        if not product.is_active:
            errors.append({"row": offset, "message": "物料已停用"})
            continue
        if parsed["code"] in seen:
            errors.append({"row": offset, "message": f"物料编号「{parsed['code']}」重复"})
            continue
        seen.add(parsed["code"])
        segment_id = None
        if parsed["department"]:
            segment_id, message = _resolve_department(parsed["department"], segments, departments)
            if message:
                errors.append({"row": offset, "message": message})
                continue
        items.append(_item(parsed, product, categories, partners, colors, units, segment_id))

    if data_rows == 0:
        raise OwnProductMaterialImportError("没有可导入的数据")
    return {"items": items, "failed": len(errors), "errors": errors[:MAX_ERRORS]}


def _item(parsed, product, categories, partners, colors, units, segment_id) -> dict:
    category = categories.get(product.category_id) if product.category_id else None
    partner = partners.get(product.partner_id)
    color = colors.get(product.color_id) if product.color_id else None
    unit = units.get(product.pricing_unit_id) if product.pricing_unit_id else None
    return {
        "qty": _num(parsed["qty"]) if parsed["qty_set"] else "1",
        "qty_set": parsed["qty_set"],
        "consume_segment_id": segment_id,
        "consume_segment_set": bool(parsed["department"]),
        "supplier_product": {
            "id": product.id,
            "product_code": product.product_code,
            "name": product.name,
            "category_id": product.category_id,
            "category_name": category.name if category else None,
            "image_url": product.image_url,
            "color_name": color.name if color else None,
            "partner_name": partner.name if partner else None,
            "pricing_unit_name": unit.name if unit else None,
            "unit_price": _num(product.unit_price),
        },
    }


def _read_rows(content: bytes) -> list[tuple]:
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except (InvalidFileException, BadZipFile, ValueError, OSError) as exc:
        raise OwnProductMaterialImportError("无法读取 Excel，请使用 .xlsx 模板") from exc
    try:
        ws = wb["产品物料"] if "产品物料" in wb.sheetnames else wb.active
        if ws is None:
            raise OwnProductMaterialImportError("Excel 里没有工作表")
        return [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _locate_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for index, raw in enumerate(rows[:8]):
        columns: dict[str, int] = {}
        for col, value in enumerate(raw):
            field = _HEADER_FIELDS.get(_norm_header(value))
            if field and field not in columns:
                columns[field] = col
        if "code" in columns:
            return index, columns
    raise OwnProductMaterialImportError("请使用产品物料导入模板，表头需包含物料编号")


def _parse_row(raw: tuple, columns: dict[str, int]) -> dict | str | None:
    def cell(field: str) -> str:
        idx = columns.get(field)
        if idx is None or idx >= len(raw):
            return ""
        return _cell_text(raw[idx])

    code = cell("code")
    qty_text = cell("qty")
    department = cell("department")
    if not any((code, qty_text, department)):
        return None
    if not code:
        return "请填写物料编号"
    if len(code) > 50:
        return "物料编号不能超过 50 字"
    qty = Decimal("1")
    qty_set = qty_text != ""
    if qty_set:
        try:
            qty = _parse_qty(qty_text)
        except OwnProductMaterialImportError as exc:
            return exc.message
    return {
        "code": code,
        "qty": qty,
        "qty_set": qty_set,
        "department": department or None,
    }


def _resolve_department(name: str, segments: list[ProcessSegment], departments: list[Department]):
    text = name.strip()
    alt = text[:-1] if text.endswith("部") and len(text) > 1 else f"{text}部"

    def active_segment(label: str) -> ProcessSegment | None:
        for seg in segments:
            if seg.is_active and (seg.name or "").strip() == label:
                return seg
        return None

    def department_segment(label: str) -> tuple[int | None, str | None]:
        for dept in departments:
            if not dept.is_active or (dept.name or "").strip() != label or not dept.process_segment_id:
                continue
            seg = next((item for item in segments if item.id == dept.process_segment_id), None)
            if seg is None:
                continue
            if not seg.is_active:
                return None, "消耗部门已停用"
            return seg.id, None
        return None, None

    hit = active_segment(text)
    if hit:
        return hit.id, None
    found_id, message = department_segment(text)
    if message or found_id:
        return found_id, message
    hit = active_segment(alt)
    if hit:
        return hit.id, None
    found_id, message = department_segment(alt)
    if message or found_id:
        return found_id, message
    for seg in segments:
        if not seg.is_active and (seg.name or "").strip() in {text, alt}:
            return None, "消耗部门已停用"
    return None, "消耗部门不存在"


def _norm_header(value) -> str:
    text = _cell_text(value).replace(" ", "").replace("\u3000", "")
    text = text.rstrip("*＊")
    for suffix in ("（必填）", "(必填)", "（选填）", "(选填)"):
        if text.endswith(suffix):
            text = text[: -len(suffix)]
    return text


def _cell_text(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def _parse_qty(value) -> Decimal:
    text = _cell_text(value).replace(",", "").replace("，", "")
    if not text:
        raise OwnProductMaterialImportError("用量不是数字")
    try:
        qty = Decimal(text)
    except InvalidOperation as exc:
        raise OwnProductMaterialImportError("用量不是数字") from exc
    if qty < 0:
        raise OwnProductMaterialImportError("用量不能为负")
    return qty.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)


def _num(value: Decimal | None) -> str | None:
    if value is None:
        return None
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"
